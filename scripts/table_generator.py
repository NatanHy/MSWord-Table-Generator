from geosphere import GeoSphere, GeoSphereInfo
from typing import Tuple, Dict
from fixed_table import FixedTable
from docx import Document
from table_config import *
from table_config import configure_document, configure_table

def span_text_over_cells(table : FixedTable, pos1 : Tuple[int, int], pos2 : Tuple[int, int], text : str):
    cell = table.cell(*pos1).merge(table.cell(*pos2))
    cell.text = text

def set_text(table : FixedTable, pos : Tuple[int, int], text : str):
    cell = table.cell(*pos)
    cell.text = text

def make_bold(table : FixedTable, pos1 : Tuple[int, int], pos2 : Tuple[int, int]):
    for i in range(pos1[0], pos2[0]):
        for j in range(pos1[1], pos2[1]):
            cell = table.cell(i, j)
            paragraphs = cell.paragraphs
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
            

def add_info(table : FixedTable, info : GeoSphereInfo, variable_descriptions : Dict[str, str]):
    variables = info.indicies(0)
    num_periods = info.num_time_periods()

    time_periods = info.indicies(2)
    row = 2

    for var in variables:
        desc = variable_descriptions[var]
        for i in range(num_periods):
            time_period = time_periods[i + 1]

            table.cell(row, 0).text = desc    
            table.cell(row, 1).text = \
                info.get_value(var, "Variable influence on process", "Influence present?", "Yes/No") + \
                "\n" + \
                info.get_value(var, "Variable influence on process", "Influence present?", "Description")
            table.cell(row, 2).text = time_period
            table.cell(row, 3).text = \
                info.get_value(var, "Variable influence on process", time_period, "Rationale")
            
            table.cell(row, 4).text = \
                info.get_value(var, "Process influence on variable", "Influence present?", "Yes/No") + \
                "\n" + \
                info.get_value(var, "Process influence on variable", "Influence present?", "Description")
            table.cell(row, 5).text = time_period
            table.cell(row, 6).text = \
                info.get_value(var, "Process influence on variable", time_period, "Rationale")
            
            row += 1

def merge_table_rows(table : FixedTable, force_cutoffs=[]):
    for col in range(table.cols):
        prev_cell = None
        start_cell_row = 0

        for row in range(table.rows):
            cell = table.cell(row, col)

            if prev_cell is not None and cell.text == prev_cell.text and not row in force_cutoffs:
                if row == table.rows - 1:
                    if row != start_cell_row:
                        text_before_merge = prev_cell.text
                        table.cell(row, col).merge(table.cell(start_cell_row, col))
                        prev_cell.text = text_before_merge
                continue
            else:
                if prev_cell is None:
                    start_cell_row = row
                    prev_cell = cell
                    continue

                if row != start_cell_row + 1:
                    text_before_merge = prev_cell.text
                    table.cell(row - 1, col).merge(table.cell(start_cell_row, col))
                    prev_cell.text = text_before_merge

                start_cell_row = row
                prev_cell = cell

def generate_table(geosphere : GeoSphere, variable_descriptions : Dict[str, str], output_filename : str):
    word_document = Document()
    configure_document(word_document)

    info = geosphere.get_info()
    num_variables = info.num_variables()
    num_periods = info.num_time_periods()

    num_rows = 2 + num_variables * num_periods
    table = FixedTable(word_document, num_rows, 7)

    # Make header
    span_text_over_cells(table, (0, 0), (1, 0), "Variables")
    span_text_over_cells(table, (0, 1), (0, 3), "Variable influence on process")    
    span_text_over_cells(table, (0, 4), (0, 6), "Process influence on variables")  

    set_text(table, (1, 1), "Influence present? (Yes/No Description)")
    set_text(table, (1, 2), "Time period/Climate domain")
    set_text(table, (1, 3), "Handling of influence \n (How/If not — Why)")
    set_text(table, (1, 4), "Influence present? (Yes/No Description)")
    set_text(table, (1, 5), "Time period/Climate domain")
    set_text(table, (1, 6), "Handling of influence \n (How/If not — Why)")

    add_info(table, info, variable_descriptions)

    num_time_periods = info.num_time_periods()
    cutoffs = [2 + i * num_time_periods for i in range(info.num_variables())]
    merge_table_rows(table, force_cutoffs=cutoffs)

    configure_table(table)
    make_bold(table, (0, 0), (2, 7))

    word_document.save(output_filename)