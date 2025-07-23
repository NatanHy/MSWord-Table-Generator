from table_generation.component import Component
from table_generation.fixed_table import FixedTable
from table_generation.parser import Parser
from typing import Tuple, Dict, List
from docx import Document
from docx.table import _Cell
import docx.document
from utils.formatting import format_raw_value, style, format_table, add_table_heading

def _get_col_sequences(table : FixedTable, col : int, force_cutoffs) -> List[Tuple[_Cell, _Cell]]:
    start = 0
    cur = 1
    start_cell = table.cell(start, col)
    seqs = []

    # Iterate over rows
    while cur < table.num_rows:
        cell = table.cell(cur, col)

        # Found a row with different text, or a forced cutoff
        if cell.text != start_cell.text or cur in force_cutoffs:
            # If more than 1 consecutive cell with identical text, merge
            if cur - start > 1:
                end_cell = table.cell(cur - 1, col)
                seqs.append((start_cell, end_cell))
            start = cur # Next sequence starts at cur
            start_cell = cell
        cur += 1

    # Run one more check to include last sequence
    if cur - start > 1:
        end_cell = table.cell(cur - 1, col)
        seqs.append((start_cell, end_cell))

    return seqs

def merge_table_rows(table : FixedTable, force_cutoffs=[]):
    """
    Merge cells with identical text in consecutive rows.
    """
    for col in range(table.num_cols):
        seqs = _get_col_sequences(table, col, force_cutoffs)
        for start_cell, end_cell in seqs:
            text_before_merge = start_cell.text
            start_cell.merge(end_cell)
            start_cell.text = text_before_merge

def generate_table_in_document(
        word_document : docx.document.Document, 
        component : Component, 
        variable_names : Dict[str, str], 
        code : str, 
        parser=Parser()
        ):
    """
    Generates a word document with a table specifying information for the given component. 
    """

    info = component.get_info()

    # Parse and execute table dsl file
    parser.parse(code)
    table_state = parser.execute(info, variable_names)

    add_table_heading(word_document, component)
    # Using fixed table class since the table shape is known after execution
    table = FixedTable(word_document, table_state.rows, table_state.cols)

    # Text needs to be added before merging
    for i in range(table_state.rows):
        for j in range(table_state.cols):
            cell = table.cell(i, j)
            text_obj = table_state.arr[i][j]
            cell.text = format_raw_value(text_obj.text)

    for span in table_state.spans:
        cell1 = table.cell(*span.pos1)
        cell2 = table.cell(*span.pos2)
        cell1.merge(cell2)
        cell1.text = span.text

    merge_table_rows(table, force_cutoffs=table_state.force_cutoffs)
    
    # Styling needs to be done after mergin
    for i in range(table_state.rows):
        for j in range(table_state.cols):
            cell = table.cell(i, j)
            text_obj = table_state.arr[i][j]

            style(cell, text_obj.style)

    # Apply table-wide configuration
    format_table(table, table_state.format)