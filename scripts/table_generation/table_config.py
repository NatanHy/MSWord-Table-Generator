import docx.document
from table_generation.fixed_table import FixedTable
from docx.shared import Cm, Pt
from typing import Tuple

# ===============================================
# Configure document by changing these variables
# ===============================================

# page margins, in centimeters
LEFT_MARGIN = 1.5
RIGHT_MARGIN = 1.5

# font size for tables, in Pts
TABLE_FONT = "Times New Roman"
TABLE_FONT_SIZE = 10
TABLE_STYLE = "Table Grid"

def make_bold(table : FixedTable, pos1 : Tuple[int, int], pos2 : Tuple[int, int]):
    for i in range(pos1[0], pos2[0]):
        for j in range(pos1[1], pos2[1]):
            cell = table.cell(i, j)
            paragraphs = cell.paragraphs
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

def configure_document(doc : docx.document.Document): 
    sections = doc.sections

    for section in sections:
        section.left_margin = Cm(LEFT_MARGIN)
        section.right_margin = Cm(RIGHT_MARGIN)

def configure_table(table : FixedTable):
    for row in range(table.rows):
        for col in range(table.cols):
            cell = table.cell(row, col)

            paragraphs = cell.paragraphs
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    run.font.name = TABLE_FONT
                    font = run.font
                    font.size = Pt(TABLE_FONT_SIZE)

    make_bold(table, (0, 0), (2, 7))

    table._table.style = TABLE_STYLE